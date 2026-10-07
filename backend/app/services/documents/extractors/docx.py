from pathlib import Path
from zipfile import BadZipFile

from docx import Document as WordDocument
from docx.opc.exceptions import PackageNotFoundError

from app.services.documents.extractors.base import (
    DocumentExtractionError,
    DocumentExtractor,
    ExtractedDocument,
    ExtractedSection,
)


class DOCXExtractor(DocumentExtractor):
    def extract(self, path: Path) -> ExtractedDocument:
        try:
            document = WordDocument(path)
            paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs]
            table_rows = [
                " | ".join(cell.text.strip() for cell in row.cells)
                for table in document.tables
                for row in table.rows
            ]
            text = "\n\n".join(part for part in paragraphs + table_rows if part)
            raw_title = document.core_properties.title
        except (OSError, ValueError, BadZipFile, PackageNotFoundError) as exc:
            raise DocumentExtractionError("The DOCX file could not be parsed") from exc
        if not text:
            raise DocumentExtractionError("The DOCX file contains no extractable text")
        title = raw_title.strip() if isinstance(raw_title, str) else None
        return ExtractedDocument((ExtractedSection(text),), title or None)
