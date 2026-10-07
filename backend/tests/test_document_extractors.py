from pathlib import Path

import pytest
from docx import Document as WordDocument
from pypdf import PdfWriter

from app.services.documents.extractors.base import DocumentExtractionError
from app.services.documents.extractors.docx import DOCXExtractor
from app.services.documents.extractors.pdf import PDFExtractor


def test_docx_extractor_reads_paragraphs_and_tables(tmp_path: Path) -> None:
    path = tmp_path / "manual.docx"
    document = WordDocument()
    document.core_properties.title = "Support Manual"
    document.add_paragraph("ResolveOps support guidance")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Refund"
    table.cell(0, 1).text = "Thirty days"
    document.save(path)

    extracted = DOCXExtractor().extract(path)

    assert extracted.title == "Support Manual"
    assert "ResolveOps support guidance" in extracted.sections[0].text
    assert "Refund | Thirty days" in extracted.sections[0].text


def test_pdf_extractor_rejects_pdf_without_extractable_text(tmp_path: Path) -> None:
    path = tmp_path / "scanned.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    with path.open("wb") as output:
        writer.write(output)

    with pytest.raises(DocumentExtractionError, match="require OCR"):
        PDFExtractor().extract(path)
